import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("SIH 2026 — Problem Statement PS146")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0A, 0x36, 0x63)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(14)
    run_sub = subtitle_p.add_run("AI-Powered Monitoring & Analysis of Bitcoin Transaction Traffic\nComprehensive Problem Analysis, Technical Architecture & Strategic Execution Plan")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    # Callout Box / Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Deployment Target", "Ubuntu 24.04 LTS (Offline / Air-Gapped Linux & WSL2)"),
        ("Dataset Specification", "v2.0 Finalized (82,078 Transactions, Dual-Layer On-Chain + P2P Telemetry)"),
        ("Primary Storage Strategy", "In-Memory Pandas DataFrames (High-Performance Zero-DB startup, Postgres fallback)"),
        ("Graph & ML Engine", "NetworkX Heterogeneous Graph + XGBoost/LightGBM with TreeSHAP Explainability")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.text = k
        c2.text = v
        set_cell_background(c1, "F0F4F8")
        set_cell_background(c2, "FAFAFA")
        set_cell_margins(c1, top=100, bottom=100, left=150, right=150)
        set_cell_margins(c2, top=100, bottom=100, left=150, right=150)
        c1.paragraphs[0].runs[0].font.bold = True
        c1.paragraphs[0].runs[0].font.size = Pt(10)
        c2.paragraphs[0].runs[0].font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1: Detailed Explanation of PS146
    h1 = doc.add_heading("1. Detailed Explanation of Problem Statement PS146", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.color.rgb = RGBColor(0x0A, 0x36, 0x63)

    p1 = doc.add_paragraph(
        "Problem Statement PS146 addresses a critical capability gap in law enforcement and financial intelligence operations: "
        "detecting, tracking, and reconstructing illicit cryptocurrency money flows across the Bitcoin blockchain. "
        "Traditional blockchain forensic tools rely solely on post-confirmation ledger heuristics (clustering heuristics, address tags), "
        "which are often defeated by modern evasion techniques such as peeling chains, multi-hop layering, CoinJoin mixing services, "
        "and rapid cross-jurisdiction routing."
    )
    p1.paragraph_format.space_after = Pt(6)

    p2 = doc.add_paragraph(
        "The core innovation of PS146 lies in Dual-Layer Fusion Forensics: combining on-chain financial accounting with pre-block "
        "peer-to-peer (P2P) network broadcast metadata. When an entity broadcasts a Bitcoin transaction, it propagates across the P2P gossip network "
        "prior to being mined into a block. By capturing the origin relay IP, Autonomous System Number (ASN), ISP infrastructure type, "
        "and propagation latency delta (timestamp - relay_timestamp), the forensic platform can identify criminal syndicates utilizing "
        "bulletproof hosting, Tor exit nodes, or VPN gateways even when on-chain amounts are obfuscated."
    )
    p2.paragraph_format.space_after = Pt(10)

    # Key Objectives Bullet List
    doc.add_heading("Core Objectives of the System:", level=2)
    objectives = [
        ("Ingest Dual-Layer Telemetry: ", "Synchronously merge UTXO ledger transactions with network relay broadcast telemetry on txid without data loss."),
        ("Multi-Entity Graph Construction: ", "Model relationships as a heterogeneous graph containing Wallet, Transaction, and IP nodes interconnected by SENT, RECEIVED, and BROADCAST edges."),
        ("Algorithmic Typology Detection: ", "Isolate distinct laundering typologies including peeling chains (change reuse), layering (fan-out/fan-in reconvergence), mixing pools (CoinJoin equal-value rounds), and ransomware campaigns."),
        ("Real Machine Learning Classification: ", "Train predictive models that output calibrated illicit risk probabilities without artificial label leakage."),
        ("Explainable Forensic Intelligence (XAI): ", "Generate court-ready, human-readable explanations powered by SHAP feature attributions detailing exactly why a candidate structure was flagged."),
        ("100% Offline Air-Gapped Operation: ", "Deliver an air-gapped, zero-cloud system capable of running in law enforcement sensitive environments on native Linux / WSL2.")
    ]
    for bold_prefix, desc in objectives:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r_b = bp.add_run(bold_prefix)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        bp.add_run(desc)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 2: Main Things Needed to Know
    h2 = doc.add_heading("2. Key Architectural Pillars & Crucial Requirements", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.color.rgb = RGBColor(0x0A, 0x36, 0x63)

    # Table of Crucial Requirements
    table2 = doc.add_table(rows=6, cols=3)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Domain / Rule", "Specification & Requirement", "Critical Detail"]
    for j, h in enumerate(headers):
        cell = table2.rows[0].cells[j]
        cell.text = h
        set_cell_background(cell, "0A3663")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    rules_data = [
        ("Zero Label Leakage Invariant", "is_illicit & pattern_type strictly EXCLUDED from features & heuristics", "Mandatory system rule. Ground-truth labels flow ONLY to offline evaluation scripts. Automated assertions prevent leakage."),
        ("Dataset Schema (v2.0)", "82,078 transactions across 17,613 scenario groups (15/15 integrity PASS)", "Multi-I/O array columns (input_addresses, etc.) parsed via json.loads(). 30 hard-negative exchange clusters planted."),
        ("Storage Architecture", "In-Memory Pandas DataFrames (Server startup load)", "Zero external DB dependencies required for hackathon runtime. PostgreSQL maintained strictly as a standby fallback."),
        ("Graph Modeling", "NetworkX heterogeneous graph (Wallet, Transaction, IP)", "No Neo4j setup overhead. NetworkX computes degree metrics, centrality, and shortest-path trace traversal in-memory."),
        ("Offline Linux Safety", "Air-gapped runnability, POSIX pathing, LF line endings", "Zero runtime CDN calls (Google Fonts bundled locally). pathlib.Path used exclusively for cross-platform Linux safety.")
    ]
    for i, (r_dom, r_spec, r_det) in enumerate(rules_data, start=1):
        row = table2.rows[i]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.text = r_dom
        c1.text = r_spec
        c2.text = r_det
        bg = "F7FAFC" if i % 2 == 1 else "FFFFFF"
        for c in (c0, c1, c2):
            set_cell_background(c, bg)
            set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            c.paragraphs[0].runs[0].font.size = Pt(9.5)
        c0.paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 3: Phased Strategic Execution Plan
    h3 = doc.add_heading("3. Strategic Roadmap & Phased Execution Plan", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    for r in h3.runs:
        r.font.color.rgb = RGBColor(0x0A, 0x36, 0x63)

    phases = [
        ("Phase 1: Environment & Repository Hardening (Immediate)", [
            "Configure repository .gitattributes with `* text=auto eol=lf` to prevent line-ending corruption.",
            "Sanitize documentation markdown links from Windows absolute `file:///c:/...` to clean relative paths.",
            "Scaffold `/backend/app/` structure: `core/`, `models/`, `api/`, `services/`, and `ingestion/`.",
            "Create `backend/requirements.txt` and launch baseline FastAPI app on `localhost:8000` with `/health` route."
        ]),
        ("Phase 2: In-Memory Data Ingestion Engine", [
            "Implement `loader.py` using `pathlib.Path` to load `blockchain_transactions.csv` + `network_metadata.csv`.",
            "Implement `parser.py` to deserialize JSON multi-I/O array columns and compute `propagation_delta_ms`.",
            "Enforce startup validation asserting 82,078 loaded rows and zero join orphans on `txid`.",
            "Hold master DataFrames in global application state for instant sub-millisecond retrieval."
        ]),
        ("Phase 3: Response Schemas & Core API Endpoints", [
            "Implement Pydantic response models in `schemas.py` matching `API_CONTRACT.md` exactly.",
            "Build `GET /entity/{address}` computing real balances, tx counts, and exchange flags from CSV data.",
            "Build `GET /transaction/{txid}` serving full dual-layer ledger + P2P network telemetry.",
            "Build `GET /graph/{scenario_id}` exporting heterogeneous Wallet, Transaction, and IP nodes/edges.",
            "Build `GET /trace?src=&dst=` running BFS pathfinding over transaction transfer chains."
        ]),
        ("Phase 4: Graph Engine & Typology Detectors (NetworkX)", [
            "Construct in-memory NetworkX directed graph from parsed ingestion DataFrames.",
            "Implement topological heuristic for peeling chain traversal (change output reuse across >=3 hops).",
            "Implement layering detector (fan-out N outputs followed by fan-in consolidation).",
            "Implement CoinJoin mixing detector (equal-denomination multi-party clusters via Louvain modularity)."
        ]),
        ("Phase 5: Machine Learning & Explainability Engine (Owner-Led)", [
            "Engineer feature matrix: financial aggregates, script type one-hots, timing delta, network infrastructure one-hots.",
            "Automate safety unit test asserting strict absence of `is_illicit`, `pattern_type`, `scenario_id`, `split` in features.",
            "Train LightGBM / XGBoost multi-class and binary models using 5-fold `GroupKFold` on `scenario_id`.",
            "Integrate TreeSHAP to compute local feature attributions and feed `/alerts` & `/alerts/{id}/evidence`."
        ]),
        ("Phase 6: Frontend Integration & Visualizer Reconciliation", [
            "Reconcile terminal CLI and Three.js 3D visualizer from Ethereum/EVM mocks to real Bitcoin UTXO schemas.",
            "Connect UI commands (`graph`, `inspect`, `trace`, `dmesg`, `alerts`) directly to live FastAPI backend.",
            "Bundle font assets locally in `public/fonts/` to ensure zero runtime network calls."
        ]),
        ("Phase 7: End-to-End Testing & Air-Gapped Packaging", [
            "Run complete scenario validation flow (e.g. `peel_001`) from raw CSV -> NetworkX -> ML -> API -> UI.",
            "Verify 100% offline air-gapped runnability on native Ubuntu 24.04 Linux environment.",
            "Prepare final presentation deck and forensic case walkthrough."
        ])
    ]

    for phase_title, tasks in phases:
        ph_head = doc.add_heading(phase_title, level=2)
        ph_head.paragraph_format.space_before = Pt(10)
        ph_head.paragraph_format.space_after = Pt(4)
        for t in tasks:
            tp = doc.add_paragraph(style='List Bullet')
            tp.paragraph_format.space_after = Pt(2)
            tp.add_run(t)

    # Save document
    output_path = os.path.abspath("SIH_PS146_Project_Plan.docx")
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_document()
